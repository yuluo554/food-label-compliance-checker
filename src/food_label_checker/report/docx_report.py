"""docx 审查报告生成（plan/04 §7）。依赖 python-docx（report extra），惰性导入。

报告结构：结论摘要 → 问题清单（分级）→ 待人工确认（独立成节）→ 证据表 →
整改建议 → 审查员签署栏 → 免责声明。LLM 参与标注（card.llm_assisted）写入
报告元信息并在免责声明中说明。

0 外链纪律：不添加任何超链接/外链图片，word/_rels/*.rels 中不得出现
TargetMode="External"（tests/test_report.py 硬断言）。
"""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List

from ..models import LEVEL_FAIL, LEVEL_MANUAL


def _import_docx():
    try:
        import docx  # noqa: F401
    except ImportError as exc:  # pragma: no cover - 依赖缺失路径（测试用 mock 注入覆盖）
        raise ImportError(
            "docx 报告需要安装 report 扩展：pip install 'food-label-compliance-checker[report]'"
        ) from exc
    return docx


def _iter_findings(result: Dict[str, Any]) -> List[Dict[str, Any]]:
    """跨规则集摊平 findings（保留 ruleset_id 于字典内）。"""
    out: List[Dict[str, Any]] = []
    for rid, res in result.get("results", {}).items():
        for finding in res.get("findings", []):
            finding = dict(finding)
            finding.setdefault("ruleset_id", rid)
            out.append(finding)
    return out


def _basis_text(finding: Dict[str, Any]) -> str:
    basis = finding.get("basis") or {}
    standard = basis.get("standard") or ""
    clause = basis.get("clause") or ""
    return f"{standard} {clause}".strip()


def _add_table(doc, headers: List[str], rows: List[List[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    try:
        table.style = "Table Grid"
    except KeyError:  # pragma: no cover - 模板差异兜底
        pass
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)


def _verdict(result: Dict[str, Any]) -> str:
    findings = _iter_findings(result)
    if any(f.get("level") == LEVEL_FAIL for f in findings):
        return "检出不合规项，需整改后复审。"
    if any(f.get("level") == LEVEL_MANUAL for f in findings):
        return "未检出确定性不合规项；存在待人工确认项，需人工复核后出具最终结论。"
    return "全部检查项合规。"


def _set_east_asian_font(doc) -> None:
    """正文中文用宋体（纯外观增强，模板差异时静默跳过）。"""
    try:
        from docx.oxml.ns import qn
        from docx.shared import Pt

        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(10.5)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    except Exception:  # noqa: BLE001 —— 外观失败不影响内容
        pass


def generate_report(result: Dict[str, Any], out_path) -> Path:
    """由 check_text/run_pipeline 的结果 JSON 生成 docx 审查报告。

    返回写入的文件路径。依赖 python-docx（report extra），缺失时抛出带
    安装提示的 ImportError（CLI 捕获后以退出码 2 优雅降级）。
    """
    _import_docx()
    from docx import Document

    card = result.get("card", {})
    fallback = result.get("fallback") or {}
    out_path = Path(out_path)
    if out_path.parent and not out_path.parent.exists():
        out_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    _set_east_asian_font(doc)

    # ---- 标题与元信息 -------------------------------------------------------
    doc.add_heading("预包装食品标签合规审查报告", level=0)
    llm_line = (
        "是（字段抽取经逐字摘录校验，数值结论均来自确定性规则）"
        if card.get("llm_assisted")
        else "否"
    )
    doc.add_paragraph(
        f"生成时间：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}\n"
        f"引擎版本：v{result.get('engine_version', '?')}    "
        f"规则集：{'、'.join(result.get('ruleset_ids', []))}\n"
        f"LLM 兜底参与：{llm_line}    兜底状态：{fallback.get('status', 'off')}"
    )

    # ---- 一、结论摘要 --------------------------------------------------------
    doc.add_heading("一、结论摘要", level=1)
    doc.add_paragraph(_verdict(result))
    stat_rows = [
        [
            rid,
            stats.get("checked", 0),
            stats.get("pass", 0),
            stats.get("fail", 0),
            stats.get("manual", 0),
        ]
        for rid, res in result.get("results", {}).items()
        for stats in [res.get("stats", {})]
    ]
    _add_table(
        doc,
        ["规则集", "检查项", "合规", "不合规", "待人工确认"],
        stat_rows,
    )
    if "dual_diff" in result:
        dd = result["dual_diff"]
        doc.add_paragraph(
            f"双标尺对照（{dd.get('baseline')} → {dd.get('target')}）："
            f"2025 新增不合规 {dd.get('new_fails', [])}；"
            f"相对放宽/消除 {dd.get('resolved_fails', [])}。"
        )
    nodes = (result.get("pipeline") or {}).get("nodes", [])
    if nodes:
        doc.add_heading("审查流程节点", level=2)
        _add_table(
            doc,
            ["节点", "状态", "耗时(ms)", "说明"],
            [[n.get("node"), n.get("status"), n.get("elapsed_ms"), n.get("detail", "")]
             for n in nodes],
        )
    for warning in result.get("warnings", []):
        doc.add_paragraph(f"警告：{warning}")

    # ---- 二、问题清单（不合规）----------------------------------------------
    doc.add_heading("二、问题清单（不合规）", level=1)
    fails = [f for f in _iter_findings(result) if f.get("level") == LEVEL_FAIL]
    if fails:
        _add_table(
            doc,
            ["规则集", "规则ID", "问题说明", "依据条款"],
            [
                [f["ruleset_id"], f["rule_id"], f.get("message", ""), _basis_text(f)]
                for f in fails
            ],
        )
    else:
        doc.add_paragraph("未发现不合规项。")

    # ---- 三、待人工确认（独立成节）-------------------------------------------
    doc.add_heading("三、待人工确认", level=1)
    manuals = [f for f in _iter_findings(result) if f.get("level") == LEVEL_MANUAL]
    if manuals:
        doc.add_paragraph(
            "以下结论由解析覆盖不足等原因转入人工复核，未经复核不得作为最终结论。"
        )
        _add_table(
            doc,
            ["规则集", "规则ID", "事项说明", "依据条款"],
            [
                [f["ruleset_id"], f["rule_id"], f.get("message", ""), _basis_text(f)]
                for f in manuals
            ],
        )
    else:
        doc.add_paragraph("无待人工确认项。")

    # ---- 四、证据表 -----------------------------------------------------------
    doc.add_heading("四、证据表", level=1)
    evidential = [
        f for f in fails + manuals if (f.get("evidence") or {}).get("quote")
    ]
    if evidential:
        _add_table(
            doc,
            ["规则ID", "区域", "原文摘录", "字符位置"],
            [
                [
                    f["rule_id"],
                    (f.get("evidence") or {}).get("region", ""),
                    (f.get("evidence") or {}).get("quote", ""),
                    str((f.get("evidence") or {}).get("span") or ""),
                ]
                for f in evidential
            ],
        )
    else:
        doc.add_paragraph("本报告的问题清单与待人工确认项均无挂接证据。")

    # ---- 五、整改建议 -----------------------------------------------------------
    doc.add_heading("五、整改建议", level=1)
    if fails:
        for i, f in enumerate(fails, start=1):
            advice = f.get("advice") or f.get("message", "")
            doc.add_paragraph(f"{i}. [{f['rule_id']}] {advice}")
    else:
        doc.add_paragraph("无。")

    # ---- 六、审查员签署栏 ---------------------------------------------------------
    doc.add_heading("六、审查员签署栏", level=1)
    _add_table(
        doc,
        ["审查员", "审查日期", "签字"],
        [["", "", ""], ["复核人", "", ""]],
    )
    doc.add_paragraph("本报告为系统自动生成初稿，须经人工复核签署后生效。")

    # ---- 七、免责声明 ---------------------------------------------------------------
    doc.add_heading("七、免责声明", level=1)
    disclaimer = (
        "本报告基于规则引擎对标签文本的自动审查结果生成，覆盖 GB 7718-2011/2025 与 "
        "GB 28050 系列既定检查项，不构成法律意见或监管结论；判定依据以现行有效标准"
        "原文及监管部门要求为准。标签原始载体（包装实物、图案、字号与版面）不在文本"
        "审查范围内。"
    )
    if card.get("llm_assisted"):
        disclaimer += (
            "本报告部分字段由 LLM 兜底抽取（原文摘录经逐字校验并回填，数值结论均来自"
            "确定性规则），相关字段已在参数卡中以 llm_assisted 标注。"
        )
    doc.add_paragraph(disclaimer)

    doc.save(str(out_path))
    return out_path
